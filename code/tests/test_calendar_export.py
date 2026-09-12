"""Official header exception, explicit dates and independent day-total audit."""
import sys
import unittest
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/"code/src"))
from cumcm_c.calendar_export import annotate_calendar, audit_calendar_comments, purchase_slot


class CalendarExportTests(unittest.TestCase):
    def test_last_column_is_current_midnight_and_calendar_total_is_once(self):
        wb=Workbook();sheet=wb.active;sheet.title="计划购电量"
        date="2025-03-01"
        for col,slot in enumerate(list(range(1,144))+[0],2):
            a,b=slot*10,(slot+1)*10
            label=f"{a//60}:{a%60:02d}-{b//60}:{b%60:02d}" if slot else "0:00-0:10+1"
            sheet.cell(1,col,label);sheet.cell(2,col,float(slot+1))
            self.assertEqual(purchase_slot(label),slot)
        sheet.cell(2,1,pd.Timestamp(date));sheet.cell(2,146,float(np.arange(1,145).sum()))
        frame=pd.DataFrame({"date":[date]*144,"slot":np.arange(144),"time_convention":"chronological_start"})
        mapping=annotate_calendar(wb,frame,[sheet.title])
        audit_calendar_comments(wb,[sheet.title])
        last=mapping.iloc[-1]
        self.assertEqual(last.delivery_start,"2025-03-01 00:00:00")
        self.assertEqual(last.source_observation_date,"2025-02-28")
        self.assertEqual(sheet.cell(1,145).value,"0:00-0:10+1")
        self.assertEqual(sheet.cell(2,145).value,1)
        self.assertEqual(mapping.delivery_start.nunique(),144)
        sheet.cell(2,145).comment=None
        with self.assertRaises(ValueError):audit_calendar_comments(wb,[sheet.title])
        annotate_calendar(wb,frame,[sheet.title]);sheet.cell(2,146).value+=1
        with self.assertRaises(ValueError):audit_calendar_comments(wb,[sheet.title])


if __name__ == "__main__":unittest.main()
