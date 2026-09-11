#!/usr/bin/env python3
"""快速核对 C 题数据附件和官方结果模板。"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "code" / "src"))

from cumcm_c.export import RESULT_FILES
from cumcm_c.paths import PATHS


def workbook_summary(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheets = ", ".join(
        f"{sheet.title}({sheet.max_row}x{sheet.max_column})" for sheet in workbook.worksheets
    )
    workbook.close()
    return sheets


def main() -> int:
    missing: list[Path] = []
    print("Data attachments:")
    for number in range(1, 5):
        path = PATHS.attachment_dir / f"附件{number}.xlsx"
        if not path.exists():
            missing.append(path)
            continue
        print(f"  OK  {path.name}: {workbook_summary(path)}")

    print("Result templates:")
    for name in RESULT_FILES:
        path = PATHS.result_template_dir / name
        if not path.exists():
            missing.append(path)
            continue
        print(f"  OK  {name}: {workbook_summary(path)}")

    if missing:
        for path in missing:
            print(f"  MISSING  {path}")
        return 1
    print("Input check: PASS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
