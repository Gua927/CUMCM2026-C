"""Official result4-2/result4-3 workbooks, read-back validation and Q4 paper tables."""
from __future__ import annotations

from copy import copy
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .paths import PATHS
from .data import time_minutes
from .models.problem3_evaluation import daily_summary, emergency_intervals

# workbook -> (template, {sheet: (per-slot energy column, daily fee column)}, strategy)
# result4-2 mirrors result2: one purchase sheet whose daily fee is the total bill.
# result4-3 mirrors result3: plan sheet with the plan charge, adjustment sheet with the total bill.
WORKBOOKS = {"result4-2.xlsx": ("result4-2.xlsx", {"计划购电量": ("planned_grid_kwh", "cost_yuan")}, "q42_forecast"),
             "result4-3.xlsx": ("result4-3.xlsx", {"计划购电量": ("planned_grid_kwh", "planned_cost_yuan"),
                                                   "调整购电量": ("adjusted_grid_kwh", "cost_yuan")}, "q43_forecast")}


def write_workbook(frame, target, template, sheets):
    target = Path(target).resolve()
    if target.exists() or PATHS.attachment_dir in target.parents:
        raise ValueError("Only a new output workbook may be written")
    wb = load_workbook(PATHS.result_template_dir/template)
    days = {date: day.reset_index(drop=True) for date, day in frame.groupby("date", sort=False)}
    daily = daily_summary(frame).set_index("date"); mapping = []
    for title, (energy, money) in sheets.items():
        sheet = wb[title]
        for row in range(2, sheet.max_row+1):
            date = str(pd.Timestamp(sheet.cell(row, 1).value).date()); day = days[date]
            for col in range(2, 146):
                label = sheet.cell(1, col).value; minute = time_minutes(label.split("-")[0]) % 1440
                item = day.iloc[minute//10]; sheet.cell(row, col, float(item[energy]))
                if row == 2 and not mapping or (row == 2 and mapping and mapping[-1]["sheet"] == title and len(mapping) < 144):
                    mapping.append({"sheet": title, "column": col, "template_label": label, "canonical_interval": item.interval, "slot": minute//10})
            sheet.cell(row, 146, float(daily.loc[date, energy])); sheet.cell(row, 147, float(daily.loc[date, money]))
    sheet = wb["充放电量"]
    prototypes = [[copy(sheet.cell(r, c)) for c in range(1, 7)] for r in range(2, 8)]
    sheet.delete_rows(2, sheet.max_row-1)
    for day_index, (date, day) in enumerate(days.items()):
        first = 2+6*day_index
        for k in range(6):
            for c in range(1, 7):
                src = prototypes[k][c-1]; dst = sheet.cell(first+k, c)
                dst._style = copy(src._style); dst.number_format = src.number_format
                if c in (2, 5):
                    dst.value = src.value
            if k == 0:
                sheet.cell(first, 1, pd.Timestamp(date).to_pydatetime())
            segment = day.iloc[k*24:(k+1)*24]
            sheet.cell(first+k, 3, float(segment.charge_kwh.sum())); sheet.cell(first+k, 4, float(segment.discharge_kwh.sum()))
        sheet.cell(first, 6, float(day.soc_start_kwh.iloc[0])); sheet.cell(first+1, 6, float(day.soc_end_kwh.iloc[-1]))
    sheet = wb["紧急购电量"]; styles = [copy(sheet.cell(2, c)._style) for c in range(1, 4)]
    sheet.delete_rows(2, sheet.max_row-1); intervals = emergency_intervals(frame); row = 2
    for date in days:
        group = intervals[intervals.date == date]
        entries = [("无", 0.0)] if group.empty else list(zip(group.interval, group.emergency_grid_kwh))
        for k, (label, amount) in enumerate(entries):
            for c in range(1, 4):
                sheet.cell(row, c)._style = copy(styles[c-1])
            if k == 0:
                sheet.cell(row, 1, pd.Timestamp(date).to_pydatetime())
            sheet.cell(row, 2, label); sheet.cell(row, 3, float(amount)); row += 1
    wb.save(target); wb.close()
    return pd.DataFrame(mapping)


def validate_workbook(target, frame, template, sheets, tolerance=2e-5):
    original = load_workbook(PATHS.result_template_dir/template)
    wb = load_workbook(target); daily = daily_summary(frame).set_index("date")
    days = {date: day.reset_index(drop=True) for date, day in frame.groupby("date", sort=False)}

    def same(a, b):
        if not isinstance(a, (int, float)) or abs(a-float(b)) > tolerance:
            raise ValueError(f"Workbook mismatch: {a} vs {b}")
    try:
        if wb.sheetnames != original.sheetnames:
            raise ValueError("Sheet names changed")
        for name in wb.sheetnames:
            a, b = wb[name], original[name]
            if a.max_column != b.max_column:
                raise ValueError("Column count changed")
            for c in range(1, b.max_column+1):
                if a.cell(1, c).value != b.cell(1, c).value:
                    raise ValueError("Header changed")
        for title, (energy, money) in sheets.items():
            sheet = wb[title]; raw = original[title]
            if sheet.max_row != 335:
                raise ValueError("Expected 334 days")
            for row in range(1, 336):
                for col in range(1, 148):
                    cell = sheet.cell(row, col); source = raw.cell(row, col)
                    if tuple(cell._style or [0]*9) != tuple(source._style or [0]*9):
                        raise ValueError("Purchase style changed")
                    if (row == 1 or col == 1) and cell.value != source.value:
                        raise ValueError("Date/time label changed")
                if row == 1:
                    continue
                date = str(pd.Timestamp(sheet.cell(row, 1).value).date()); day = days[date]
                for col in range(2, 146):
                    slot = time_minutes(sheet.cell(1, col).value.split("-")[0]) % 1440//10
                    same(sheet.cell(row, col).value, day[energy].iloc[slot])
                same(sheet.cell(row, 146).value, daily.loc[date, energy]); same(sheet.cell(row, 147).value, daily.loc[date, money])
        sheet = wb["充放电量"]
        if sheet.max_row != 1+6*len(days):
            raise ValueError("Storage dates incomplete")
        for i, (date, day) in enumerate(days.items()):
            row = 2+i*6
            if pd.Timestamp(sheet.cell(row, 1).value) != pd.Timestamp(date):
                raise ValueError("Storage date mismatch")
            for k in range(6):
                for col in range(1, 7):
                    if tuple(sheet.cell(row+k, col)._style or [0]*9) != tuple(original[sheet.title].cell(k+2, col)._style or [0]*9):
                        raise ValueError("Storage style mismatch")
                for col in (2, 5):
                    if sheet.cell(row+k, col).value != original[sheet.title].cell(k+2, col).value:
                        raise ValueError("Storage interval label mismatch")
                same(sheet.cell(row+k, 3).value, day.charge_kwh.iloc[k*24:(k+1)*24].sum())
                same(sheet.cell(row+k, 4).value, day.discharge_kwh.iloc[k*24:(k+1)*24].sum())
            same(sheet.cell(row, 6).value, day.soc_start_kwh.iloc[0]); same(sheet.cell(row+1, 6).value, day.soc_end_kwh.iloc[-1])
        expected = emergency_intervals(frame); actual = []; date = None
        for row in wb["紧急购电量"].iter_rows(min_row=2, values_only=True):
            if row[0] is not None:
                date = str(pd.Timestamp(row[0]).date())
            if row[1] == "无":
                same(row[2], 0); continue
            actual.append((date, row[1], row[2]))
        if len(actual) != len(expected):
            raise ValueError("Emergency intervals count mismatch")
        for got, want in zip(actual, expected.itertuples()):
            if got[:2] != (want.date, want.interval):
                raise ValueError("Emergency label mismatch")
            same(got[2], want.emergency_grid_kwh)
        same(sum(x[2] for x in actual), frame.emergency_grid_kwh.sum())
        return {"passed": True, "template": template, "days": len(days), "purchase_sheets": list(sheets),
                "purchase_cells": len(sheets)*len(days)*144, "storage_blocks": len(days)*6, "emergency_intervals": len(actual),
                "fee_columns": {k: v[1] for k, v in sheets.items()}, "settlement": "realized attachment 4 price of each delivery slot"}
    finally:
        wb.close(); original.close()


def _purchase_table(frame, dates, label, caption, adjusted):
    daily = daily_summary(frame).set_index("date")
    head = r"时间段 & 计划/调整 & 时间段 & 计划/调整 & 时间段 & 计划/调整 \\" if adjusted else r"时间段 & kWh & 时间段 & kWh & 时间段 & kWh \\"
    lines = [r"\begin{table}[htbp]\centering\small", f"\\caption{{{caption}}}", f"\\label{{{label}}}\\begin{{tabular}}{{lrlrlr}}\\toprule", head]
    for date in dates:
        day = frame[frame.date == date].reset_index(drop=True); stats = daily.loc[date]
        lines += [r"\midrule", f"\\multicolumn{{6}}{{c}}{{{date}}}"+r" \\"]
        for hours in ((10, 12, 14), (16, 18, 20)):
            if adjusted:
                lines.append(" & ".join(f"{h}:00--{h}:10 & {day.planned_grid_kwh.iloc[h*6]:.3f}/{day.adjusted_grid_kwh.iloc[h*6]:.3f}" for h in hours)+r" \\")
            else:
                lines.append(" & ".join(f"{h}:00--{h}:10 & {day.planned_grid_kwh.iloc[h*6]:.3f}" for h in hours)+r" \\")
        if adjusted:
            lines.append(f"\\multicolumn{{3}}{{c}}{{全天计划/调整后：{stats.planned_grid_kwh:.3f}/{stats.adjusted_grid_kwh:.3f} kWh}} & "
                         + f"\\multicolumn{{3}}{{c}}{{全天购电费：{stats.cost_yuan:.2f} 元（含调整{stats.adjustment_cost_yuan:+.2f}、紧急{stats.emergency_cost_yuan:.2f}）}}"+r" \\")
        else:
            lines.append(f"\\multicolumn{{3}}{{c}}{{计划购电：{stats.planned_grid_kwh:.3f} kWh}} & "+f"\\multicolumn{{3}}{{c}}{{总费用：{stats.cost_yuan:.2f} 元（含紧急{stats.emergency_cost_yuan:.2f}）}}"+r" \\")
    lines += [r"\bottomrule\end{tabular}\end{table}"]
    return lines


def _storage_table(frame, dates, label, caption):
    lines = [r"\begin{table}[htbp]\centering\small", f"\\caption{{{caption}}}", f"\\label{{{label}}}\\begin{{tabular}}{{lrrlrr}}\\toprule",
             r"时间段 & 充电 & 放电 & 时间段 & 充电 & 放电 \\"]
    for date in dates:
        day = frame[frame.date == date].reset_index(drop=True)
        lines += [r"\midrule", f"\\multicolumn{{6}}{{c}}{{{date}}}"+r" \\"]
        for k in (0, 2, 4):
            items = []
            for j in (k, k+1):
                part = day.iloc[j*24:(j+1)*24]
                items.append(f"{j*4}:00--{(j+1)*4}:00 & {part.charge_kwh.sum():.3f} & {part.discharge_kwh.sum():.3f}")
            lines.append(" & ".join(items)+r" \\")
        lines.append(f"\\multicolumn{{3}}{{c}}{{0:00 电量：{day.soc_start_kwh.iloc[0]:.3f}}} & "+f"\\multicolumn{{3}}{{c}}{{24:00 电量：{day.soc_end_kwh.iloc[-1]:.3f}}}"+r" \\")
    lines += [r"\bottomrule\end{tabular}\end{table}"]
    return lines


def _emergency_table(emergencies, dates, label, caption):
    groups = [emergencies[emergencies.date == date].reset_index(drop=True) for date in dates]
    lines = [r"\begingroup\scriptsize", r"\setlength{\tabcolsep}{3pt}", r"\begin{longtable}{lrlrlrlr}",
             f"\\caption{{{caption}}}\\label{{{label}}}\\\\\\toprule",
             " & ".join(f"\\multicolumn{{2}}{{c}}{{{date}}}" for date in dates)+r" \\",
             " & ".join(["时间段 & kWh"]*4)+r" \\\midrule\endfirsthead", " & ".join(["时间段 & kWh"]*4)+r" \\\midrule\endhead"]
    for i in range(max(1, max(map(len, groups)))):
        items = []
        for group in groups:
            items.append(f"{group.interval.iloc[i]} & {group.emergency_grid_kwh.iloc[i]:.3f}" if i < len(group) else "-- & --")
        lines.append(" & ".join(items)+r" \\")
    lines += [r"\midrule", " & ".join(f"合计 & {g.emergency_grid_kwh.sum():.3f}" for g in groups)+r" \\", r"\bottomrule\end{longtable}\endgroup"]
    return lines


LABELS = {"q42_forecast": "4-2主策略：价格预测", "q42_fixed_profile": "4-2：按附件1价格轮廓规划", "q42_perfect_price": "4-2：当日价格预知（启发式对照）",
          "q43_forecast": "4-3主策略：价格预测+日内修正", "q43_legacy_chain": "4-3：冻结旧版信息链", "q43_cross_day_only": "4-3：仅接入跨日新预报",
          "q43_matched_bias_only": "4-3：仅同构残差校准", "q43_no_intraday": "4-3：不作日内价格修正",
          "q43_fixed_profile": "4-3：按附件1价格轮廓规划", "q43_perfect_price": "4-3：当日价格预知（启发式对照）", "q43_plan_only": "4-3：仅0:00计划"}


def tex_tables(frames, comparison, dates, target, price_table, errors, ci, check, references):
    def write(name, lines):
        (target/name).write_text("\n".join(lines)+"\n", encoding="utf-8")
    f42, f43 = frames["q42_forecast"], frames["q43_forecast"]
    write("q42_purchase.tex", _purchase_table(f42, dates, "tab:q42-purchase", "波动电价下问题二指定日期计划购电量与实际费用（表1格式，按实时价结算）", False))
    write("q42_storage.tex", _storage_table(f42, dates, "tab:q42-storage", "波动电价下问题二指定日期充放电量及日初日末储电量（表2格式，kWh）"))
    write("q42_emergency.tex", _emergency_table(emergency_intervals(f42), dates, "tab:q42-emergency", "波动电价下问题二指定日期紧急购电区间及电量（表3格式）"))
    write("q43_purchase.tex", _purchase_table(f43, dates, "tab:q43-purchase", "波动电价下问题三指定日期购电量（表1格式；每格为0:00计划量/最终合同量，kWh）", True))
    write("q43_storage.tex", _storage_table(f43, dates, "tab:q43-storage", "波动电价下问题三指定日期充放电量及日初日末储电量（表2格式，kWh）"))
    write("q43_emergency.tex", _emergency_table(emergency_intervals(f43), dates, "tab:q43-emergency", "波动电价下问题三指定日期最终紧急购电区间及电量（表3格式）"))
    lines = [r"\begin{table}[htbp]\centering\small", r"\setlength{\tabcolsep}{4pt}",
             r"\caption{波动电价下二月至十二月的实际费用（万元）与紧急购电（kWh）；所有账单按交付时段实时价结算}",
             r"\label{tab:q4-comparison}\begin{tabular}{lrrrrrr}\toprule",
             r"策略 & 计划费 & 调整费 & 紧急费 & 总费用 & 修正费用 & 紧急量 \\\midrule"]
    for row in comparison.itertuples():
        lines.append(f"{LABELS[row.strategy]} & {row.planned_cost_yuan/1e4:.3f} & {row.adjustment_cost_yuan/1e4:+.3f} & {row.emergency_cost_yuan/1e4:.3f} & "
                     f"{row.cost_yuan/1e4:.3f} & {row.value_adjusted_cost_yuan/1e4:.3f} & {row.emergency_grid_kwh/1e3:.1f}k"+r" \\")
    lines += [r"\midrule", f"固定电价参考：第二问主策略 & \\multicolumn{{5}}{{r}}{{{references['q2_fixed_price_cost']/1e4:.3f}}} & {references['q2_fixed_price_emergency']/1e3:.1f}k"+r" \\",
              f"固定电价参考：第三问主策略 & \\multicolumn{{5}}{{r}}{{{references['q3_fixed_price_cost']/1e4:.3f}}} & {references['q3_fixed_price_emergency']/1e3:.1f}k"+r" \\",
              r"\bottomrule\end{tabular}\end{table}"]
    write("q4_comparison.tex", lines)
    method_labels = {"fixed": "附件1轮廓", "lag7": "上周同日", "weekday": "两周同日均值", "regression": "净负荷回归", "combined": "融合（主模型）"}
    lines = [r"\begin{table}[htbp]\centering\small", r"\caption{实时电价日前预测的平均绝对误差（元/kWh）：一月检验与二月至十二月}",
             r"\label{tab:q4-price}\begin{tabular}{lrrrr}\toprule", r"预测方法 & 一月：当日 & 一月：次日 & 二月至十二月：当日 & 二月至十二月：次日 \\\midrule"]
    chk = check.set_index("method"); err = errors.set_index("method")
    for method in ("fixed", "lag7", "weekday", "regression", "combined"):
        lines.append(f"{method_labels[method]} & {chk.loc[method,'lead1_mae']:.4f} & {chk.loc[method,'lead2_mae']:.4f} & {err.loc[method,'lead1_mae']:.4f} & {err.loc[method,'lead2_mae']:.4f}"+r" \\")
    lines += [r"\bottomrule\end{tabular}\end{table}"]
    write("q4_price.tex", lines)
    c = comparison.set_index("strategy"); m42 = c.loc["q42_forecast"]; m43 = c.loc["q43_forecast"]
    def diff(a, b, column="cost_yuan"):
        return float(c.loc[a, column]-c.loc[b, column])
    numbers = {"QFourTwoTotalCost": f"{m42.cost_yuan/1e4:.3f}", "QFourTwoPlanCost": f"{m42.planned_cost_yuan/1e4:.3f}",
               "QFourTwoEmergencyCost": f"{m42.emergency_cost_yuan/1e4:.3f}", "QFourTwoEmergencyEnergy": f"{m42.emergency_grid_kwh:.3f}",
               "QFourTwoTerminal": f"{m42.terminal_kwh:.3f}", "QFourTwoAdjusted": f"{m42.value_adjusted_cost_yuan:.2f}", "QFourTwoCash": f"{m42.cost_yuan:.2f}",
               "QFourTwoFixedCost": f"{c.loc['q42_fixed_profile','cost_yuan']:.2f}", "QFourTwoFixedDiff": f"{diff('q42_fixed_profile','q42_forecast'):.2f}",
               "QFourTwoFixedPercent": f"{diff('q42_fixed_profile','q42_forecast')/c.loc['q42_fixed_profile','cost_yuan']*100:.2f}",
               "QFourTwoPerfectCost": f"{c.loc['q42_perfect_price','cost_yuan']:.2f}", "QFourTwoPerfectGap": f"{diff('q42_forecast','q42_perfect_price'):.2f}",
               "QFourTwoQTwoReference": f"{references['q2_fixed_price_cost']:.2f}", "QFourTwoVsQTwo": f"{m42.cost_yuan-references['q2_fixed_price_cost']:.2f}",
               "QFourTwoFixedEmergency": f"{c.loc['q42_fixed_profile','emergency_grid_kwh']:.3f}",
               "QFourTwoVsQTwoPercent": f"{(m42.cost_yuan-references['q2_fixed_price_cost'])/references['q2_fixed_price_cost']*100:.2f}",
               "QFourTwoSettlementEffect": f"{c.loc['q42_fixed_profile','cost_yuan']-references['q2_fixed_price_cost']:.2f}",
               "QFourTwoSettlementPercent": f"{(c.loc['q42_fixed_profile','cost_yuan']-references['q2_fixed_price_cost'])/references['q2_fixed_price_cost']*100:.2f}",
               "QFourThreeSettlementEffect": f"{c.loc['q43_fixed_profile','cost_yuan']-references['q3_fixed_price_cost']:.2f}",
               "QFourThreeSettlementPercent": f"{(c.loc['q43_fixed_profile','cost_yuan']-references['q3_fixed_price_cost'])/references['q3_fixed_price_cost']*100:.2f}",
               "QFourThreeVsQThreePercent": f"{(m43.cost_yuan-references['q3_fixed_price_cost'])/references['q3_fixed_price_cost']*100:.2f}",
               "QFourThreePerfectGapPercent": f"{diff('q43_forecast','q43_perfect_price')/m43.cost_yuan*100:.2f}",
               "QFourTwoPerfectGapPercent": f"{diff('q42_forecast','q42_perfect_price')/m42.cost_yuan*100:.2f}",
               "QFourThreeEmergencyDays": str(int(m43.emergency_days)), "QFourThreeAdjustedSlots": str(int(m43.adjusted_slots)),
               "QFourTwoCvar": f"{m42.emergency_cost_cvar95_yuan:.2f}", "QFourThreePlanOnlyEmergency": f"{c.loc['q43_plan_only','emergency_grid_kwh']:.3f}",
               "QFourThreeTotalCost": f"{m43.cost_yuan/1e4:.3f}", "QFourThreePlanCost": f"{m43.planned_cost_yuan/1e4:.3f}",
               "QFourThreeAdjustCost": f"{m43.adjustment_cost_yuan/1e4:+.3f}", "QFourThreeEmergencyCost": f"{m43.emergency_cost_yuan/1e4:.3f}",
               "QFourThreeEmergencyEnergy": f"{m43.emergency_grid_kwh:.3f}", "QFourThreeUpEnergy": f"{m43.adjust_up_kwh:.3f}",
               "QFourThreeDownEnergy": f"{m43.adjust_down_kwh:.3f}", "QFourThreeTerminal": f"{m43.terminal_kwh:.3f}", "QFourThreeCash": f"{m43.cost_yuan:.2f}",
               "QFourThreeAdjusted": f"{m43.value_adjusted_cost_yuan:.2f}", "QFourThreeCvar": f"{m43.emergency_cost_cvar95_yuan:.2f}",
               "QFourThreePlanOnlyCost": f"{c.loc['q43_plan_only','cost_yuan']:.2f}", "QFourThreeSavingPlanOnly": f"{diff('q43_plan_only','q43_forecast'):.2f}",
               "QFourThreeSavingPlanOnlyPercent": f"{diff('q43_plan_only','q43_forecast')/c.loc['q43_plan_only','cost_yuan']*100:.2f}",
               "QFourThreeNoIntradayCost": f"{c.loc['q43_no_intraday','cost_yuan']:.2f}", "QFourThreeNoIntradayDiff": f"{diff('q43_no_intraday','q43_forecast'):.2f}",
               "QFourThreeFixedCost": f"{c.loc['q43_fixed_profile','cost_yuan']:.2f}", "QFourThreeFixedDiff": f"{diff('q43_fixed_profile','q43_forecast'):.2f}",
               "QFourThreeFixedPercent": f"{diff('q43_fixed_profile','q43_forecast')/c.loc['q43_fixed_profile','cost_yuan']*100:.2f}",
               "QFourThreePerfectCost": f"{c.loc['q43_perfect_price','cost_yuan']:.2f}", "QFourThreePerfectGap": f"{diff('q43_forecast','q43_perfect_price'):.2f}",
               "QFourThreeQThreeReference": f"{references['q3_fixed_price_cost']:.2f}", "QFourThreeVsQThree": f"{m43.cost_yuan-references['q3_fixed_price_cost']:.2f}",
               "QFourThreeVsTwo": f"{diff('q42_forecast','q43_forecast'):.2f}", "QFourThreeVsTwoPercent": f"{diff('q42_forecast','q43_forecast')/m42.cost_yuan*100:.2f}",
               "QFourPriceMean": f"{price_table['realized_mean'].mean():.4f}", "QFourPriceDailyMin": f"{price_table['realized_mean'].min():.3f}",
               "QFourPriceDailyMax": f"{price_table['realized_mean'].max():.3f}", "QFourPriceWeekendMean": f"{price_table.loc[price_table.low_group,'realized_mean'].mean():.3f}",
               "QFourPriceOtherMean": f"{price_table.loc[~price_table.low_group,'realized_mean'].mean():.3f}",
               "QFourDeficitPriceRatio": f"{references['deficit_price_ratio']:.3f}", "QFourPriceNetCorr": f"{references['price_net_corr']:.3f}",
               "QFourPriceRSquared": f"{references['price_net_r2']:.3f}",
               # R10: the 4-2 -> 4-3 difference splits into the 0:00 external-forecast fusion (no adjustment) and the paid intraday adjustment.
               "QFourValueExternalPlan": f"{diff('q42_forecast','q43_plan_only'):.2f}", "QFourValueAdjust": f"{diff('q43_plan_only','q43_forecast'):.2f}",
               "QFourValueExternalPlanPercent": f"{diff('q42_forecast','q43_plan_only')/m42.cost_yuan*100:.2f}",
               "QFourValueAdjustPercent": f"{diff('q43_plan_only','q43_forecast')/m42.cost_yuan*100:.2f}",
               "QFourThreePerfectEmergency": f"{c.loc['q43_perfect_price','emergency_grid_kwh']:.3f}",
               "QFourTwoPerfectEmergency": f"{c.loc['q42_perfect_price','emergency_grid_kwh']:.3f}",
               # Emergency-energy reductions quoted in prose: 4-2 -> 4-3 and plan_only -> 4-3 (the adjustment step).
               "QFourEmergencyDropPercent": f"{(1-m43.emergency_grid_kwh/m42.emergency_grid_kwh)*100:.1f}",
               "QFourAdjustEmergencyDropPercent": f"{(1-m43.emergency_grid_kwh/c.loc['q43_plan_only','emergency_grid_kwh'])*100:.1f}",
               "QFourAdjustShareOfSaving": f"{diff('q43_plan_only','q43_forecast')/diff('q42_forecast','q43_forecast')*100:.1f}",
               "QFourStrategyCount": str(len(comparison))}
    for strategy, key in (("q43_legacy_chain", "Legacy"), ("q43_cross_day_only", "CrossDay"), ("q43_matched_bias_only", "MatchedBias")):
        if strategy in c.index:
            r = c.loc[strategy]
            numbers.update({f"QFourThree{key}Cost": f"{r.cost_yuan:.2f}", f"QFourThree{key}Diff": f"{diff(strategy,'q43_forecast'):.2f}",
                            f"QFourThree{key}Emergency": f"{r.emergency_grid_kwh:.3f}", f"QFourThree{key}Cvar": f"{r.emergency_cost_cvar95_yuan:.2f}",
                            f"QFourThree{key}PlanDiff": f"{diff(strategy,'q43_forecast','planned_cost_yuan'):.2f}",
                            f"QFourThree{key}AdjustDiff": f"{diff(strategy,'q43_forecast','adjustment_cost_yuan'):.2f}",
                            f"QFourThree{key}EmergencyDiff": f"{diff(strategy,'q43_forecast','emergency_cost_yuan'):.2f}"})
    if "q43_cross_day_only" in c.index and "q43_legacy_chain" in c.index:
        numbers["QFourThreeCrossDayVsLegacy"] = f"{diff('q43_cross_day_only','q43_legacy_chain'):.2f}"
    for method, key in (("fixed", "Fixed"), ("lag7", "LagSeven"), ("weekday", "Weekday"), ("regression", "Regression"), ("combined", "Combined")):
        numbers[f"QFourMAE{key}"] = f"{err.loc[method,'lead1_mae']:.4f}"; numbers[f"QFourMAE{key}Two"] = f"{err.loc[method,'lead2_mae']:.4f}"
        numbers[f"QFourCheck{key}"] = f"{chk.loc[method,'lead1_mae']:.4f}"
    for t, key in ((36, "Six"), (72, "Twelve"), (108, "Eighteen")):
        numbers[f"QFourRemaining{key}"] = f"{references['intraday'][f'remaining_{t}']:.4f}"; numbers[f"QFourStale{key}"] = f"{references['intraday'][f'stale_{t}']:.4f}"
    suffix = {"q42_fixed_profile": "TwoFixed", "q42_perfect_price": "TwoPerfect", "q43_no_intraday": "ThreeNoIntraday", "q43_fixed_profile": "ThreeFixed",
              "q43_perfect_price": "ThreePerfect", "q43_plan_only": "ThreePlanOnly", "q42_forecast": "ThreeVsTwo", "q43_legacy_chain": "ThreeLegacy",
              "q43_cross_day_only": "ThreeCrossDay", "q43_matched_bias_only": "ThreeMatchedBias"}
    for strategy, values in ci.items():
        if strategy in suffix:
            numbers[f"QFourCI{suffix[strategy]}Low"] = f"{values['ci95_mean_daily_yuan'][0]:.2f}"
            numbers[f"QFourCI{suffix[strategy]}High"] = f"{values['ci95_mean_daily_yuan'][1]:.2f}"
            numbers[f"QFourCI{suffix[strategy]}Mean"] = f"{values['mean_daily_saving_yuan']:.2f}"
    write("q4_numbers.tex", [f"\\providecommand{{\\{k}}}{{{v}}}" for k, v in numbers.items()])
