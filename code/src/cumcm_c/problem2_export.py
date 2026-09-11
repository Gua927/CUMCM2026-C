"""Official Q2 template expansion and independently verified numeric export."""
from __future__ import annotations

from copy import copy
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .paths import PATHS
from .data import time_minutes
from .models.problem2_evaluation import daily_summary, emergency_intervals


def write_workbook(frame,target):
    target=Path(target).resolve()
    if target.exists() or PATHS.attachment_dir in target.parents:
        raise ValueError("Only a new output workbook may be written")
    wb=load_workbook(PATHS.result_template_dir/"result2.xlsx");sheet=wb["计划购电量"]
    days={date:day.reset_index(drop=True) for date,day in frame.groupby("date",sort=False)}
    daily=daily_summary(frame).set_index("date");mapping=[]
    for row in range(2,sheet.max_row+1):
        date=str(pd.Timestamp(sheet.cell(row,1).value).date());day=days[date]
        for col in range(2,146):
            label=sheet.cell(1,col).value
            minute=time_minutes(label.split("-")[0])%1440
            item=day.iloc[minute//10]
            sheet.cell(row,col,float(item.planned_grid_kwh))
            if row==2:mapping.append({"column":col,"template_label":label,"canonical_interval":item.interval,"slot":minute//10})
        sheet.cell(row,146,float(daily.loc[date,"planned_grid_kwh"]))
        # The daily cash bill includes emergency purchases, separately itemized elsewhere.
        sheet.cell(row,147,float(daily.loc[date,"cost_yuan"]))
    sheet=wb["充放电量"]
    # Expand the template's ellipsis into the same six-row date block.
    prototypes=[[copy(sheet.cell(r,c)) for c in range(1,7)] for r in range(2,8)]
    sheet.delete_rows(2,sheet.max_row-1)
    for day_index,(date,day) in enumerate(days.items()):
        first=2+6*day_index
        for k in range(6):
            for c in range(1,7):
                src=prototypes[k][c-1];dst=sheet.cell(first+k,c)
                dst._style=copy(src._style);dst.number_format=src.number_format
                if c in (2,5):dst.value=src.value
            if k==0:sheet.cell(first,1,pd.Timestamp(date).to_pydatetime())
            segment=day.iloc[k*24:(k+1)*24]
            sheet.cell(first+k,3,float(segment.charge_kwh.sum()))
            sheet.cell(first+k,4,float(segment.discharge_kwh.sum()))
        sheet.cell(first,6,float(day.soc_start_kwh.iloc[0]))
        sheet.cell(first+1,6,float(day.soc_end_kwh.iloc[-1]))
    sheet=wb["紧急购电量"];styles=[copy(sheet.cell(2,c)._style) for c in range(1,4)]
    sheet.delete_rows(2,sheet.max_row-1);intervals=emergency_intervals(frame)
    row=2
    for date in days:
        group=intervals[intervals.date==date]
        if group.empty:entries=[("无",0.0)]
        else:entries=list(zip(group.interval,group.emergency_grid_kwh))
        for k,(label,amount) in enumerate(entries):
            for c in range(1,4):sheet.cell(row,c)._style=copy(styles[c-1])
            if k==0:sheet.cell(row,1,pd.Timestamp(date).to_pydatetime())
            sheet.cell(row,2,label);sheet.cell(row,3,float(amount));row+=1
    wb.save(target);wb.close()
    return pd.DataFrame(mapping)


def validate_workbook(target,frame,tolerance=2e-5):
    original=load_workbook(PATHS.result_template_dir/"result2.xlsx")
    wb=load_workbook(target);daily=daily_summary(frame).set_index("date")
    days={date:day.reset_index(drop=True) for date,day in frame.groupby("date",sort=False)}
    def same(a,b):
        if not isinstance(a,(int,float)) or abs(a-float(b))>tolerance:raise ValueError(f"Workbook mismatch: {a} vs {b}")
    try:
        if wb.sheetnames!=original.sheetnames:raise ValueError("Sheet names changed")
        for name in wb.sheetnames:
            a,b=wb[name],original[name]
            if a.max_column!=b.max_column:raise ValueError("Column count changed")
            for c in range(1,b.max_column+1):
                if a.cell(1,c).value!=b.cell(1,c).value:raise ValueError("Header changed")
        sheet=wb["计划购电量"];raw=original[sheet.title]
        if sheet.max_row!=335:raise ValueError("Expected 334 days")
        for row in range(1,336):
            for col in range(1,148):
                cell=sheet.cell(row,col);source=raw.cell(row,col)
                if tuple(cell._style or [0]*9)!=tuple(source._style or [0]*9):raise ValueError("Purchase style changed")
                if (row==1 or col==1) and cell.value!=source.value:raise ValueError("Date/time label changed")
            if row==1:continue
            date=str(pd.Timestamp(sheet.cell(row,1).value).date());day=days[date]
            for col in range(2,146):
                slot=time_minutes(sheet.cell(1,col).value.split("-")[0])%1440//10
                same(sheet.cell(row,col).value,day.planned_grid_kwh.iloc[slot])
            same(sheet.cell(row,146).value,daily.loc[date,"planned_grid_kwh"])
            same(sheet.cell(row,147).value,daily.loc[date,"cost_yuan"])
        sheet=wb["充放电量"]
        if sheet.max_row!=1+6*len(days):raise ValueError("Storage dates incomplete")
        for i,(date,day) in enumerate(days.items()):
            row=2+i*6
            if pd.Timestamp(sheet.cell(row,1).value)!=pd.Timestamp(date):raise ValueError("Storage date mismatch")
            for k in range(6):
                for col in range(1,7):
                    if tuple(sheet.cell(row+k,col)._style or [0]*9)!=tuple(original[sheet.title].cell(k+2,col)._style or [0]*9):raise ValueError("Storage style mismatch")
                for col in (2,5):
                    if sheet.cell(row+k,col).value!=original[sheet.title].cell(k+2,col).value:raise ValueError("Storage interval label mismatch")
                same(sheet.cell(row+k,3).value,day.charge_kwh.iloc[k*24:(k+1)*24].sum())
                same(sheet.cell(row+k,4).value,day.discharge_kwh.iloc[k*24:(k+1)*24].sum())
            same(sheet.cell(row,6).value,day.soc_start_kwh.iloc[0]);same(sheet.cell(row+1,6).value,day.soc_end_kwh.iloc[-1])
        expected=emergency_intervals(frame);actual=[];date=None
        for row in wb["紧急购电量"].iter_rows(min_row=2,values_only=True):
            if row[0] is not None:date=str(pd.Timestamp(row[0]).date())
            if row[1]=="无":same(row[2],0);continue
            actual.append((date,row[1],row[2]))
        if len(actual)!=len(expected):raise ValueError("Emergency intervals count mismatch")
        for got,want in zip(actual,expected.itertuples()):
            if got[:2]!=(want.date,want.interval):raise ValueError("Emergency label mismatch")
            same(got[2],want.emergency_grid_kwh)
        same(sum(x[2] for x in actual),frame.emergency_grid_kwh.sum())
        return {"passed":True,"days":len(days),"purchase_cells":len(days)*144,
                "storage_blocks":len(days)*6,"emergency_intervals":len(actual),
                "daily_cost_column":"planned_cost + emergency_cost",
                "template":"Original headers, purchase labels and dates preserved; ellipsis blocks expanded using template styles."}
    finally:wb.close();original.close()


def tex_tables(frame,comparison,emergencies,dates,target):
    def write(name,lines):(target/name).write_text("\n".join(lines)+"\n",encoding="utf-8")
    daily=daily_summary(frame).set_index("date")
    lines=[r"\begin{table}[htbp]\centering\small",r"\caption{问题二指定日期计划购电量与实际总费用（表1格式）}",
      r"\label{tab:q2-purchase}\begin{tabular}{lrlrlr}\toprule",
      r"时间段 & kWh & 时间段 & kWh & 时间段 & kWh \\"]
    for date in dates:
        day=frame[frame.date==date].reset_index(drop=True);stats=daily.loc[date]
        lines += [r"\midrule",f"\\multicolumn{{6}}{{c}}{{{date}}}"+r" \\"]
        for hours in ((10,12,14),(16,18,20)):
            lines.append(" & ".join(f"{h}:00--{h}:10 & {day.planned_grid_kwh.iloc[h*6]:.3f}" for h in hours)+r" \\")
        lines.append(f"\\multicolumn{{3}}{{c}}{{计划购电：{stats.planned_grid_kwh:.3f} kWh}} & "+f"\\multicolumn{{3}}{{c}}{{总费用：{stats.cost_yuan:.2f} 元}}"+r" \\")
    lines += [r"\bottomrule\end{tabular}\end{table}"];write("q2_purchase.tex",lines)
    lines=[r"\begin{table}[htbp]\centering\small",r"\caption{问题二指定日期充放电量及日初日末储电量（表2格式，单位：kWh）}",
      r"\label{tab:q2-storage}\begin{tabular}{lrrlrr}\toprule",
      r"时间段 & 充电 & 放电 & 时间段 & 充电 & 放电 \\"]
    for date in dates:
        day=frame[frame.date==date].reset_index(drop=True)
        lines += [r"\midrule",f"\\multicolumn{{6}}{{c}}{{{date}}}"+r" \\"]
        for k in (0,2,4):
            items=[]
            for j in (k,k+1):
                part=day.iloc[j*24:(j+1)*24]
                items.append(f"{j*4}:00--{(j+1)*4}:00 & {part.charge_kwh.sum():.3f} & {part.discharge_kwh.sum():.3f}")
            lines.append(" & ".join(items)+r" \\")
        lines.append(f"\\multicolumn{{3}}{{c}}{{0:00 电量：{day.soc_start_kwh.iloc[0]:.3f}}} & "+f"\\multicolumn{{3}}{{c}}{{24:00 电量：{day.soc_end_kwh.iloc[-1]:.3f}}}"+r" \\")
    lines += [r"\bottomrule\end{tabular}\end{table}"];write("q2_storage.tex",lines)
    groups=[emergencies[emergencies.date==date].reset_index(drop=True) for date in dates]
    lines=[r"\begingroup\scriptsize",r"\setlength{\tabcolsep}{3pt}",r"\begin{longtable}{lrlrlrlr}",
      r"\caption{问题二指定日期紧急购电区间及电量（表3格式）}\label{tab:q2-emergency}\\\toprule",
      " & ".join(f"\\multicolumn{{2}}{{c}}{{{date}}}" for date in dates)+r" \\",
      " & ".join(["时间段 & kWh"]*4)+r" \\\midrule\endfirsthead",
      " & ".join(["时间段 & kWh"]*4)+r" \\\midrule\endhead"]
    for i in range(max(map(len,groups))):
        items=[]
        for group in groups:
            items.append(f"{group.interval.iloc[i]} & {group.emergency_grid_kwh.iloc[i]:.3f}" if i<len(group) else "-- & --")
        lines.append(" & ".join(items)+r" \\")
    lines += [r"\midrule", " & ".join(f"合计 & {g.emergency_grid_kwh.sum():.3f}" for g in groups)+r" \\",r"\bottomrule\end{longtable}\endgroup"]
    write("q2_emergency.tex",lines)
    labels={"risk_policy":"一月选定风险策略","point_policy":"点预测+同一控制器","risk_mpc":"风险分位数+滚动控制","no_storage":"无储能分位数","risk_greedy":"风险分位数+贪心","risk_reference":"风险分位数+参考轨迹"}
    lines=[r"\begin{table}[htbp]\centering\small",r"\caption{二月至十二月因果策略的实际费用比较}",
       r"\label{tab:q2-comparison}\begin{tabular}{lrrrr}\toprule",
       r"策略 & 计划费/万元 & 紧急费/万元 & 总费用/万元 & 期末电量/kWh \\\midrule"]
    for row in comparison.itertuples():
        lines.append(f"{labels[row.strategy]} & {row.planned_cost_yuan/1e4:.3f} & {row.emergency_cost_yuan/1e4:.3f} & {row.cost_yuan/1e4:.3f} & {row.terminal_kwh:.1f}"+r" \\")
    lines += [r"\bottomrule\end{tabular}\end{table}"];write("q2_comparison.tex",lines)
    main=comparison.iloc[0];base=comparison.set_index("strategy").loc["no_storage"]
    numbers={"QTwoTotalCost":f"{main.cost_yuan/1e4:.3f}","QTwoPlanCost":f"{main.planned_cost_yuan/1e4:.3f}",
       "QTwoEmergencyCost":f"{main.emergency_cost_yuan/1e4:.3f}","QTwoEmergencyEnergy":f"{main.emergency_grid_kwh:.3f}",
       "QTwoSaving":f"{(1-main.cost_yuan/base.cost_yuan)*100:.2f}","QTwoTerminal":f"{main.terminal_kwh:.3f}",
       "QTwoLoadMAE":f"{main.load_mae_kw:.2f}","QTwoPVMAE":f"{main.pv_mae_kw:.2f}",
       "QTwoCoverage":f"{main.risk_pointwise_coverage*100:.2f}"}
    write("q2_numbers.tex",[f"\\providecommand{{\\{k}}}{{{v}}}" for k,v in numbers.items()])
