"""Official result3 template expansion, read-back validation and paper tables."""
from __future__ import annotations

from copy import copy
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .paths import PATHS
from .data import time_minutes
from .calendar_export import purchase_slot, annotate_calendar, audit_calendar_comments
from .models.problem3_evaluation import daily_summary, emergency_intervals

# Sheet -> (per-slot energy column, daily fee column). The plan sheet carries the 0:00 plan
# and its charge; the adjustment sheet carries the last effective contract and the full
# daily bill (plan + adjustment settlement + emergency), i.e. the problem's total cost.
PURCHASE_SHEETS = {"计划购电量": ("planned_grid_kwh", "planned_cost_yuan"), "调整购电量": ("adjusted_grid_kwh", "cost_yuan")}


def write_workbook(frame, target):
    target = Path(target).resolve()
    if target.exists() or PATHS.attachment_dir in target.parents:
        raise ValueError("Only a new output workbook may be written")
    wb = load_workbook(PATHS.result_template_dir/"result3.xlsx")
    days = {date: day.reset_index(drop=True) for date, day in frame.groupby("date", sort=False)}
    daily = daily_summary(frame).set_index("date"); mapping = []
    for title, (energy, money) in PURCHASE_SHEETS.items():
        sheet = wb[title]
        for row in range(2, sheet.max_row+1):
            date = str(pd.Timestamp(sheet.cell(row, 1).value).date()); day = days[date]
            for col in range(2, 146):
                label = sheet.cell(1, col).value; minute = purchase_slot(label)*10
                item = day.iloc[minute//10]; sheet.cell(row, col, float(item[energy]))
                if row == 2 and title == "计划购电量":
                    mapping.append({"column": col, "template_label": label, "canonical_interval": item.interval, "slot": minute//10})
            sheet.cell(row, 146, float(daily.loc[date, energy])); sheet.cell(row, 147, float(daily.loc[date, money]))
    sheet = wb["充放电量"]
    prototypes = [[copy(sheet.cell(r, c)) for c in range(1, 7)] for r in range(2, 8)]
    sheet.delete_rows(2, sheet.max_row-1)
    for day_index, (date, day) in enumerate(days.items()):
        first = 2+6*day_index
        for k in range(6):
            for c in range(1, 7):
                src = prototypes[k][c-1]; dst = sheet.cell(first+k, c)
                dst._style = copy(src._style); dst.number_format = src.number_format
                if c in (2, 5):
                    dst.value = src.value
            if k == 0:
                sheet.cell(first, 1, pd.Timestamp(date).to_pydatetime())
            segment = day.iloc[k*24:(k+1)*24]
            sheet.cell(first+k, 3, float(segment.charge_kwh.sum())); sheet.cell(first+k, 4, float(segment.discharge_kwh.sum()))
        sheet.cell(first, 6, float(day.soc_start_kwh.iloc[0])); sheet.cell(first+1, 6, float(day.soc_end_kwh.iloc[-1]))
    sheet = wb["紧急购电量"]; styles = [copy(sheet.cell(2, c)._style) for c in range(1, 4)]
    sheet.delete_rows(2, sheet.max_row-1); intervals = emergency_intervals(frame); row = 2
    for date in days:
        group = intervals[intervals.date == date]
        entries = [("无", 0.0)] if group.empty else list(zip(group.interval, group.emergency_grid_kwh))
        for k, (label, amount) in enumerate(entries):
            for c in range(1, 4):
                sheet.cell(row, c)._style = copy(styles[c-1])
            if k == 0:
                sheet.cell(row, 1, pd.Timestamp(date).to_pydatetime())
            sheet.cell(row, 2, label); sheet.cell(row, 3, float(amount)); row += 1
    mapping = annotate_calendar(wb, frame, list(PURCHASE_SHEETS))
    wb.save(target); wb.close()
    return mapping


def validate_workbook(target, frame, tolerance=2e-5):
    original = load_workbook(PATHS.result_template_dir/"result3.xlsx")
    wb = load_workbook(target); daily = daily_summary(frame).set_index("date")
    days = {date: day.reset_index(drop=True) for date, day in frame.groupby("date", sort=False)}

    def same(a, b):
        if not isinstance(a, (int, float)) or abs(a-float(b)) > tolerance:
            raise ValueError(f"Workbook mismatch: {a} vs {b}")
    try:
        if "time_convention" in frame:
            audit_calendar_comments(wb, list(PURCHASE_SHEETS))
        if wb.sheetnames != original.sheetnames:
            raise ValueError("Sheet names changed")
        for name in wb.sheetnames:
            a, b = wb[name], original[name]
            if a.max_column != b.max_column:
                raise ValueError("Column count changed")
            for c in range(1, b.max_column+1):
                if a.cell(1, c).value != b.cell(1, c).value:
                    raise ValueError("Header changed")
        for title, (energy, money) in PURCHASE_SHEETS.items():
            sheet = wb[title]; raw = original[title]
            if sheet.max_row != 335:
                raise ValueError("Expected 334 days")
            for row in range(1, 336):
                for col in range(1, 148):
                    cell = sheet.cell(row, col); source = raw.cell(row, col)
                    if tuple(cell._style or [0]*9) != tuple(source._style or [0]*9):
                        raise ValueError("Purchase style changed")
                    if (row == 1 or col == 1) and cell.value != source.value:
                        raise ValueError("Date/time label changed")
                if row == 1:
                    continue
                date = str(pd.Timestamp(sheet.cell(row, 1).value).date()); day = days[date]
                for col in range(2, 146):
                    slot = time_minutes(sheet.cell(1, col).value.split("-")[0]) % 1440//10
                    same(sheet.cell(row, col).value, day[energy].iloc[slot])
                same(sheet.cell(row, 146).value, daily.loc[date, energy]); same(sheet.cell(row, 147).value, daily.loc[date, money])
        sheet = wb["充放电量"]
        if sheet.max_row != 1+6*len(days):
            raise ValueError("Storage dates incomplete")
        for i, (date, day) in enumerate(days.items()):
            row = 2+i*6
            if pd.Timestamp(sheet.cell(row, 1).value) != pd.Timestamp(date):
                raise ValueError("Storage date mismatch")
            for k in range(6):
                for col in range(1, 7):
                    if tuple(sheet.cell(row+k, col)._style or [0]*9) != tuple(original[sheet.title].cell(k+2, col)._style or [0]*9):
                        raise ValueError("Storage style mismatch")
                for col in (2, 5):
                    if sheet.cell(row+k, col).value != original[sheet.title].cell(k+2, col).value:
                        raise ValueError("Storage interval label mismatch")
                same(sheet.cell(row+k, 3).value, day.charge_kwh.iloc[k*24:(k+1)*24].sum())
                same(sheet.cell(row+k, 4).value, day.discharge_kwh.iloc[k*24:(k+1)*24].sum())
            same(sheet.cell(row, 6).value, day.soc_start_kwh.iloc[0]); same(sheet.cell(row+1, 6).value, day.soc_end_kwh.iloc[-1])
        expected = emergency_intervals(frame); actual = []; date = None
        for row in wb["紧急购电量"].iter_rows(min_row=2, values_only=True):
            if row[0] is not None:
                date = str(pd.Timestamp(row[0]).date())
            if row[1] == "无":
                same(row[2], 0); continue
            actual.append((date, row[1], row[2]))
        if len(actual) != len(expected):
            raise ValueError("Emergency intervals count mismatch")
        for got, want in zip(actual, expected.itertuples()):
            if got[:2] != (want.date, want.interval):
                raise ValueError("Emergency label mismatch")
            same(got[2], want.emergency_grid_kwh)
        same(sum(x[2] for x in actual), frame.emergency_grid_kwh.sum())
        return {"passed": True, "days": len(days), "purchase_cells": 2*len(days)*144, "storage_blocks": len(days)*6,
                "emergency_intervals": len(actual),
                "plan_sheet_fee": "sum of price x 0:00 plan", "adjustment_sheet_fee": "plan charge + adjustment settlement + emergency (total daily bill)",
                "template": "Original headers, purchase labels and dates preserved; ellipsis blocks expanded using template styles."}
    finally:
        wb.close(); original.close()


LABELS = {"adjust_all": "主策略：三时点调整（新信息链）", "legacy_chain": "冻结旧版信息链", "cross_day_only": "仅接入跨日新预报",
          "matched_bias_only": "仅同构残差校准", "plan_only": "仅0:00计划", "adjust_6": "仅6:00", "adjust_12": "仅12:00", "adjust_18": "仅18:00",
          "adjust_6_12": "6:00+12:00", "adjust_6_18": "6:00+18:00", "adjust_12_18": "12:00+18:00",
          "stale_forecast": "三时点调整，预报冻结", "external_only": "三时点调整，仅附件3", "own_forecast_plan_only": "第二问策略（无附件3）",
          "own_forecast_adjust_all": "无附件3，三时点调整", "penalty_rule": "主策略，不退款口径", "symmetric_rule": "主策略，单一分位数"}
SEGMENTS = {"today_remaining": "当日剩余", "tomorrow_in_range": "次日：版本覆盖范围内", "tomorrow_beyond_range": "次日：覆盖范围外"}


def tex_tables(frame, comparison, emergencies, dates, target, diagnostics, ci, check, descriptions,
               horizon=None, horizon_legacy=None, certificates=None, audits=None, alignment=None):
    def write(name, lines):
        (target/name).write_text("\n".join(lines)+"\n", encoding="utf-8")
    daily = daily_summary(frame).set_index("date")
    lines = [r"\begin{table}[htbp]\centering\small", r"\caption{问题三指定日期购电量（表1格式；每格为0:00计划量/最终计划量，单位：kWh）}",
             r"\label{tab:q3-purchase}\begin{tabular}{lrlrlr}\toprule", r"时间段 & 计划/调整 & 时间段 & 计划/调整 & 时间段 & 计划/调整 \\"]
    for date in dates:
        day = frame[frame.date == date].reset_index(drop=True); stats = daily.loc[date]
        lines += [r"\midrule", f"\\multicolumn{{6}}{{c}}{{{date}}}"+r" \\"]
        for hours in ((10, 12, 14), (16, 18, 20)):
            lines.append(" & ".join(f"{h}:00--{h}:10 & {day.planned_grid_kwh.iloc[h*6]:.3f}/{day.adjusted_grid_kwh.iloc[h*6]:.3f}" for h in hours)+r" \\")
        lines.append(f"\\multicolumn{{3}}{{c}}{{全天计划/调整后：{stats.planned_grid_kwh:.3f}/{stats.adjusted_grid_kwh:.3f} kWh}} & "
                     + f"\\multicolumn{{3}}{{c}}{{全天购电费：{stats.cost_yuan:.2f} 元（含调整{stats.adjustment_cost_yuan:+.2f}、紧急{stats.emergency_cost_yuan:.2f}）}}"+r" \\")
    lines += [r"\bottomrule\end{tabular}\end{table}"]; write("q3_purchase.tex", lines)
    lines = [r"\begin{table}[htbp]\centering\small", r"\caption{问题三指定日期充放电量及日初日末储电量（表2格式，单位：kWh）}",
             r"\label{tab:q3-storage}\begin{tabular}{lrrlrr}\toprule", r"时间段 & 充电 & 放电 & 时间段 & 充电 & 放电 \\"]
    for date in dates:
        day = frame[frame.date == date].reset_index(drop=True)
        lines += [r"\midrule", f"\\multicolumn{{6}}{{c}}{{{date}}}"+r" \\"]
        for k in (0, 2, 4):
            items = []
            for j in (k, k+1):
                part = day.iloc[j*24:(j+1)*24]
                items.append(f"{j*4}:00--{(j+1)*4}:00 & {part.charge_kwh.sum():.3f} & {part.discharge_kwh.sum():.3f}")
            lines.append(" & ".join(items)+r" \\")
        lines.append(f"\\multicolumn{{3}}{{c}}{{0:00 电量：{day.soc_start_kwh.iloc[0]:.3f}}} & "+f"\\multicolumn{{3}}{{c}}{{24:00 电量：{day.soc_end_kwh.iloc[-1]:.3f}}}"+r" \\")
    lines += [r"\bottomrule\end{tabular}\end{table}"]; write("q3_storage.tex", lines)
    groups = [emergencies[emergencies.date == date].reset_index(drop=True) for date in dates]
    lines = [r"\begingroup\scriptsize", r"\setlength{\tabcolsep}{3pt}", r"\begin{longtable}{lrlrlrlr}",
             r"\caption{问题三指定日期最终紧急购电区间及电量（表3格式）}\label{tab:q3-emergency}\\\toprule",
             " & ".join(f"\\multicolumn{{2}}{{c}}{{{date}}}" for date in dates)+r" \\",
             " & ".join(["时间段 & kWh"]*4)+r" \\\midrule\endfirsthead", " & ".join(["时间段 & kWh"]*4)+r" \\\midrule\endhead"]
    for i in range(max(1, max(map(len, groups)))):
        items = []
        for group in groups:
            items.append(f"{group.interval.iloc[i]} & {group.emergency_grid_kwh.iloc[i]:.3f}" if i < len(group) else "-- & --")
        lines.append(" & ".join(items)+r" \\")
    lines += [r"\midrule", " & ".join(f"合计 & {g.emergency_grid_kwh.sum():.3f}" for g in groups)+r" \\", r"\bottomrule\end{longtable}\endgroup"]
    write("q3_emergency.tex", lines)
    labels = LABELS
    chain = [s for s in ("adjust_all", "legacy_chain", "cross_day_only", "matched_bias_only") if s in set(comparison.strategy)]
    lines = [r"\begin{table}[htbp]\centering\small", r"\setlength{\tabcolsep}{4pt}",
             r"\caption{二月至十二月各策略的实际费用与调整量（万元、kWh）；除结算口径行外均按退款口径结算，除信息链对照外均采用新信息链}",
             r"\label{tab:q3-comparison}\begin{tabular}{lrrrrrrr}\toprule",
             r"策略 & 计划费 & 调整费 & 紧急费 & 总费用 & 修正费用 & 上调/下调量 & 紧急量 \\\midrule"]
    for row in comparison.itertuples():
        if row.strategy in chain[1:]:
            continue
        lines.append(f"{labels[row.strategy]} & {row.planned_cost_yuan/1e4:.3f} & {row.adjustment_cost_yuan/1e4:+.3f} & {row.emergency_cost_yuan/1e4:.3f} & "
                     f"{row.cost_yuan/1e4:.3f} & {row.value_adjusted_cost_yuan/1e4:.3f} & {row.adjust_up_kwh/1e3:.1f}/{row.adjust_down_kwh/1e3:.1f}k & {row.emergency_grid_kwh/1e3:.1f}k"+r" \\")
    lines += [r"\bottomrule\end{tabular}\end{table}"]; write("q3_comparison.tex", lines)
    if len(chain) == 4:
        c0 = comparison.set_index("strategy")
        lines = [r"\begin{table}[htbp]\centering\footnotesize", r"\setlength{\tabcolsep}{3pt}",
                 r"\caption{信息链修复的四组同参对照（二月至十二月；相同输入、初态、控制器与结算；期末电量均为8837 kWh）：现金费用及分解（万元）、紧急购电（kWh）、日紧急费用CVaR95（元）与主策略相对该变体日均节省的七日块配对95\%区间（元）}",
                 r"\label{tab:q3-information-chain}\begin{tabular}{lccrrrrrrl}\toprule",
                 r"变体 & 跨日新预报 & 同构残差 & 总费用 & 计划费 & 调整费 & 紧急费 & 紧急量 & CVaR95 & 日均节省区间 \\\midrule"]
        for s in chain:
            r = c0.loc[s]; flag = lambda v: "是" if bool(v) else "否"
            interval = "--" if s == "adjust_all" or s not in ci else f"[{ci[s]['ci95_mean_daily_yuan'][0]:.1f}, {ci[s]['ci95_mean_daily_yuan'][1]:.1f}]"
            lines.append(f"{labels[s]} & {flag(r.cross_day_release)} & {flag(r.matched_bias)} & {r.cost_yuan/1e4:.3f} & {r.planned_cost_yuan/1e4:.3f} & {r.adjustment_cost_yuan/1e4:+.3f} & "
                         f"{r.emergency_cost_yuan/1e4:.3f} & {r.emergency_grid_kwh/1e3:.1f}k & {r.emergency_cost_cvar95_yuan:.0f} & {interval}"+r" \\")
        lines += [r"\bottomrule\end{tabular}\end{table}"]; write("q3_information_chain.tex", lines)
    if horizon is not None and horizon_legacy is not None:
        lines = [r"\begin{table}[htbp]\centering\small", r"\setlength{\tabcolsep}{4pt}",
                 r"\caption{调整时刻风险轮廓按发布时刻与提前期分段的实际覆盖率（净负荷不超过轮廓的时段比例，\%）与光伏点预报MAE（kW）：新信息链与冻结旧版在同一分段上对照；次日按该发布版本的24小时覆盖范围分段（新信息链只在范围内使用该版本），名义覆盖率为0.7/0.9（当日）与0.8（次日）}",
                 r"\label{tab:q3-horizon}\begin{tabular}{llrrrrrr}\toprule",
                 r"发布 & 分段 & 时段数 & 新：0.7/0.8覆盖 & 新：0.9覆盖 & 新：光伏MAE & 旧：0.7/0.8覆盖 & 旧：光伏MAE \\\midrule"]
        h1 = horizon.set_index(["release_slot", "segment"]); h0 = horizon_legacy.set_index(["release_slot", "segment"])
        for (release, segment), r in h1.iterrows():
            old = h0.loc[(release, segment)]
            high = f"{r.high_coverage*100:.1f}" if segment == "today_remaining" else "--"
            lines.append(f"{r.release} & {SEGMENTS[segment]} & {int(r.rows)} & {r.low_coverage*100:.1f} & {high} & {r.pv_mae_kw:.1f} & {old.low_coverage*100:.1f} & {old.pv_mae_kw:.1f}"+r" \\")
        lines += [r"\bottomrule\end{tabular}\end{table}"]; write("q3_horizon.tex", lines)
    lines = [r"\begin{table}[htbp]\centering\small", r"\caption{光伏预报平均绝对误差（kW，二月至十二月）：发布后24小时窗口及当日剩余时段}",
             r"\label{tab:q3-forecast}\begin{tabular}{lrrrrr}\toprule",
             r"发布时刻 & 附件3 & 本地趋势 & 融合 & 当日剩余：0:00版本 & 当日剩余：最新版本 \\\midrule"]
    for row in diagnostics.itertuples():
        same = "--" if np.isnan(row.stale_same_day_mae_kw) else f"{row.stale_same_day_mae_kw:.1f}"
        fresh = "--" if np.isnan(row.fresh_same_day_mae_kw) else f"{row.fresh_same_day_mae_kw:.1f}"
        lines.append(f"{row.release} & {row.external_mae_kw:.1f} & {row.own_mae_kw:.1f} & {row.combined_mae_kw:.1f} & {same} & {fresh}"+r" \\")
    lines += [r"\bottomrule\end{tabular}\end{table}"]; write("q3_forecast.tex", lines)
    c = comparison.set_index("strategy"); m = c.loc["adjust_all"]
    def diff(a, b, column="cost_yuan"):
        return float(c.loc[a, column]-c.loc[b, column])
    numbers = {"QThreeTotalCost": f"{m.cost_yuan/1e4:.3f}", "QThreePlanCost": f"{m.planned_cost_yuan/1e4:.3f}",
               "QThreeAdjustCost": f"{m.adjustment_cost_yuan/1e4:+.3f}", "QThreeUpCost": f"{m.upward_cost_yuan/1e4:.3f}",
               "QThreeDownCost": f"{m.downward_cost_yuan/1e4:+.3f}", "QThreeEmergencyCost": f"{m.emergency_cost_yuan/1e4:.3f}",
               "QThreeEmergencyEnergy": f"{m.emergency_grid_kwh:.3f}", "QThreeUpEnergy": f"{m.adjust_up_kwh:.3f}",
               "QThreeDownEnergy": f"{m.adjust_down_kwh:.3f}", "QThreeAdjustedSlots": str(int(m.adjusted_slots)),
               "QThreeAdjustedDays": str(int(m.adjusted_days)), "QThreeTerminal": f"{m.terminal_kwh:.3f}",
               "QThreeAdjusted": f"{m.value_adjusted_cost_yuan:.2f}", "QThreeCvar": f"{m.emergency_cost_cvar95_yuan:.2f}",
               "QThreeEmergencyDays": str(int(m.emergency_days)), "QThreeEmergencySlots": str(int(m.emergency_slots)),
               "QThreeCash": f"{m.cost_yuan:.2f}",
               "QThreePlanOnlyCost": f"{c.loc['plan_only','cost_yuan']:.2f}", "QThreeSavingPlanOnly": f"{diff('plan_only','adjust_all'):.2f}",
               "QThreeSavingPlanOnlyPercent": f"{diff('plan_only','adjust_all')/c.loc['plan_only','cost_yuan']*100:.2f}",
               "QThreeQTwoCost": f"{c.loc['own_forecast_plan_only','cost_yuan']:.2f}",
               "QThreeSavingQTwo": f"{diff('own_forecast_plan_only','adjust_all'):.2f}",
               "QThreeSavingQTwoPercent": f"{diff('own_forecast_plan_only','adjust_all')/c.loc['own_forecast_plan_only','cost_yuan']*100:.2f}",
               "QThreeValueExternalPlan": f"{diff('own_forecast_plan_only','plan_only'):.2f}",
               "QThreeValueAdjustNoInfo": f"{diff('plan_only','stale_forecast'):.2f}",
               "QThreeValueFreshInfo": f"{diff('stale_forecast','adjust_all'):.2f}",
               "QThreeStaleCost": f"{c.loc['stale_forecast','cost_yuan']:.2f}",
               "QThreeMarginalSix": f"{diff('adjust_12_18','adjust_all'):.2f}", "QThreeMarginalTwelve": f"{diff('adjust_6_18','adjust_all'):.2f}",
               "QThreeMarginalEighteen": f"{diff('adjust_6_12','adjust_all'):.2f}",
               "QThreeSingleSix": f"{diff('plan_only','adjust_6'):.2f}", "QThreeSingleTwelve": f"{diff('plan_only','adjust_12'):.2f}",
               "QThreeSingleEighteen": f"{diff('plan_only','adjust_18'):.2f}",
               "QThreeExternalOnlyCost": f"{c.loc['external_only','cost_yuan']:.2f}", "QThreeExternalOnlyDiff": f"{diff('external_only','adjust_all'):.2f}",
               "QThreeOwnAdjustCost": f"{c.loc['own_forecast_adjust_all','cost_yuan']:.2f}",
               "QThreePenaltyCost": f"{c.loc['penalty_rule','cost_yuan']:.2f}", "QThreePenaltyDiff": f"{diff('penalty_rule','adjust_all'):.2f}",
               "QThreePenaltyDown": f"{c.loc['penalty_rule','adjust_down_kwh']:.3f}",
               "QThreeSymmetricCost": f"{c.loc['symmetric_rule','cost_yuan']:.2f}", "QThreeSymmetricDiff": f"{diff('symmetric_rule','adjust_all'):.2f}",
               "QThreePlanOnlyEmergency": f"{c.loc['plan_only','emergency_grid_kwh']:.3f}",
               "QThreePlanOnlyCvar": f"{c.loc['plan_only','emergency_cost_cvar95_yuan']:.2f}",
               "QThreeAdjustedDiffPlanOnly": f"{diff('plan_only','adjust_all','value_adjusted_cost_yuan'):.2f}",
               "QThreeAdjustedDiffQTwo": f"{diff('own_forecast_plan_only','adjust_all','value_adjusted_cost_yuan'):.2f}",
               "QThreeCoverage": f"{m.risk_pointwise_coverage*100:.2f}", "QThreePVMAE": f"{m.pv_mae_kw:.2f}", "QThreeLatestPVMAE": f"{m.latest_pv_mae_kw:.2f}",
               "QThreeLoadMAE": f"{m.load_mae_kw:.2f}"}
    d = diagnostics.set_index("release")
    for release, key in (("0:00", "Zero"), ("6:00", "Six"), ("12:00", "Twelve"), ("18:00", "Eighteen")):
        numbers[f"QThreeExtMAE{key}"] = f"{d.loc[release,'external_mae_kw']:.1f}"
        numbers[f"QThreeOwnMAE{key}"] = f"{d.loc[release,'own_mae_kw']:.1f}"
        numbers[f"QThreeMixMAE{key}"] = f"{d.loc[release,'combined_mae_kw']:.1f}"
        if release != "0:00":
            numbers[f"QThreeStaleMAE{key}"] = f"{d.loc[release,'stale_same_day_mae_kw']:.1f}"
            numbers[f"QThreeFreshMAE{key}"] = f"{d.loc[release,'fresh_same_day_mae_kw']:.1f}"
            if "next_day_local_mae_kw" in d:
                numbers[f"QThreeNextLocal{key}"] = f"{d.loc[release,'next_day_local_mae_kw']:.1f}"
                numbers[f"QThreeNextAware{key}"] = f"{d.loc[release,'next_day_release_aware_mae_kw']:.1f}"
    suffix = {"plan_only": "PlanOnly", "stale_forecast": "Stale", "adjust_12_18": "DropSix", "adjust_6_18": "DropTwelve",
              "adjust_6_12": "DropEighteen", "own_forecast_plan_only": "QTwo", "external_only": "ExternalOnly", "penalty_rule": "Penalty",
              "symmetric_rule": "Symmetric", "legacy_chain": "Legacy", "cross_day_only": "CrossDay", "matched_bias_only": "MatchedBias"}
    for strategy, values in ci.items():
        if strategy in suffix:
            numbers[f"QThreeCI{suffix[strategy]}Low"] = f"{values['ci95_mean_daily_yuan'][0]:.2f}"
            numbers[f"QThreeCI{suffix[strategy]}High"] = f"{values['ci95_mean_daily_yuan'][1]:.2f}"
            numbers[f"QThreeCI{suffix[strategy]}Mean"] = f"{values['mean_daily_saving_yuan']:.2f}"
    k = check.set_index("variant")
    for variant, key in (("band_derived", "Band"), ("symmetric_plan_quantile", "Symmetric"), ("band_wide", "Wide"), ("plan_only", "PlanOnly"), ("legacy_chain", "Legacy")):
        if variant in k.index:
            numbers[f"QThreeCheck{key}"] = f"{k.loc[variant,'check_adjusted_cost_yuan']:.2f}"
    for strategy, key in (("legacy_chain", "Legacy"), ("cross_day_only", "CrossDay"), ("matched_bias_only", "MatchedBias")):
        if strategy in c.index:
            r = c.loc[strategy]
            numbers[f"QThree{key}Cost"] = f"{r.cost_yuan:.2f}"; numbers[f"QThree{key}Diff"] = f"{diff(strategy, 'adjust_all'):.2f}"
            numbers[f"QThree{key}PlanDiff"] = f"{diff(strategy, 'adjust_all', 'planned_cost_yuan'):.2f}"
            numbers[f"QThree{key}AdjustDiff"] = f"{diff(strategy, 'adjust_all', 'adjustment_cost_yuan'):.2f}"
            numbers[f"QThree{key}EmergencyDiff"] = f"{diff(strategy, 'adjust_all', 'emergency_cost_yuan'):.2f}"
            numbers[f"QThree{key}Emergency"] = f"{r.emergency_grid_kwh:.3f}"; numbers[f"QThree{key}Cvar"] = f"{r.emergency_cost_cvar95_yuan:.2f}"
            numbers[f"QThree{key}Terminal"] = f"{r.terminal_kwh:.3f}"; numbers[f"QThree{key}AdjustedDiff"] = f"{diff(strategy, 'adjust_all', 'value_adjusted_cost_yuan'):.2f}"
            numbers[f"QThree{key}Up"] = f"{r.adjust_up_kwh:.3f}"; numbers[f"QThree{key}Down"] = f"{r.adjust_down_kwh:.3f}"
    if certificates is not None:
        # LP counts: the annual trajectories of the compared strategies, and the whole run including the January check.
        numbers["QThreeLPSolves"] = str(int(comparison.lp_solves.sum())) if "lp_solves" in comparison else str(int(certificates["solves"]))
        numbers["QThreeLPSolvesTotal"] = str(int(certificates["solves"]))
        numbers["QThreeLPSolvesJanuary"] = str(int(certificates["solves"])-int(comparison.lp_solves.sum())) if "lp_solves" in comparison else "0"
        numbers["QThreeSolverFailures"] = str(int(certificates.get("failures", 0)))
        numbers["QThreeDualGap"] = f"{certificates['max_primal_dual_gap_yuan']:.1e}".replace("e-", r"\times10^{-").replace("e+", r"\times10^{")+"}"
    if audits is not None and "issued_adjustments" in audits:
        numbers["QThreeAdjustmentRecords"] = str(int(audits["issued_adjustments"]["records"]))
        numbers["QThreeStrategyCount"] = str(len(comparison))
    if alignment is not None:
        a = alignment.set_index(["release_slot", "window"])
        for window, key in (("all_days", ""), ("evaluation", "Eval")):
            if (0, window) in a.index:
                r = a.loc[(0, window)]
                numbers.update({f"QThreeAlign{key}Start": f"{r.start_slot_mae_kw:.0f}", f"QThreeAlign{key}End": f"{r.end_slot_mae_kw:.0f}",
                                f"QThreeAlign{key}PrevMean": f"{r.previous_hour_mean_mae_kw:.0f}", f"QThreeAlign{key}NextMean": f"{r.next_hour_mean_mae_kw:.0f}",
                                f"QThreeAlign{key}Days": str(int(r.days))})
    if "cross_day_only" in c.index and "legacy_chain" in c.index:
        numbers["QThreeCrossDayVsLegacy"] = f"{diff('cross_day_only', 'legacy_chain'):.2f}"
    for table, prefix in ((horizon, "QThreeCov"), (horizon_legacy, "QThreeLegacyCov")):
        if table is None:
            continue
        h = table.set_index(["release_slot", "segment"])
        for release, key in ((36, "Six"), (72, "Twelve"), (108, "Eighteen")):
            for segment, seg in (("today_remaining", "Today"), ("tomorrow_in_range", "Covered"), ("tomorrow_beyond_range", "Uncovered")):
                if (release, segment) in h.index:
                    numbers[f"{prefix}{key}{seg}Low"] = f"{h.loc[(release, segment), 'low_coverage']*100:.1f}"
                    numbers[f"{prefix}{key}{seg}High"] = f"{h.loc[(release, segment), 'high_coverage']*100:.1f}"
                    numbers[f"{prefix.replace('Cov', 'PV')}{key}{seg}"] = f"{h.loc[(release, segment), 'pv_mae_kw']:.1f}"
    write("q3_numbers.tex", [f"\\providecommand{{\\{name}}}{{{value}}}" for name, value in numbers.items()])
