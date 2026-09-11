"""官方结果工作簿的安全准备与后续写回入口。"""

from __future__ import annotations

import shutil
from pathlib import Path

from .paths import PATHS


RESULT_FILES = (
    "result1.xlsx",
    "result2.xlsx",
    "result3.xlsx",
    "result4-2.xlsx",
    "result4-3.xlsx",
)


def prepare_result_workbook(name: str, *, overwrite: bool = False) -> Path:
    """复制一个官方空白模板到输出区，不修改题目附件。"""
    if name not in RESULT_FILES:
        raise ValueError(f"未知结果文件：{name}")
    source = PATHS.result_template_dir / name
    target = PATHS.result_dir / name
    if not source.exists():
        raise FileNotFoundError(source)
    PATHS.result_dir.mkdir(parents=True, exist_ok=True)
    if target.exists() and not overwrite:
        raise FileExistsError(f"输出已存在；如需覆盖请显式指定 overwrite=True：{target}")
    shutil.copy2(source, target)
    return target


def write_problem1_workbook(schedule, target: Path):
    """Write only numeric template cells. Return the explicit cyclic interval map."""
    import pandas as pd
    from openpyxl import load_workbook
    from .data import time_minutes
    from .models.evaluation import reporting_tables

    target=Path(target).resolve()
    if PATHS.attachment_dir.resolve() in target.parents:
        raise ValueError("Original attachments are read-only")
    if target.exists():
        raise FileExistsError(target)
    wb=load_workbook(PATHS.result_template_dir / "result1.xlsx")
    sheet=wb["计划购电量"]
    by_start=schedule.set_index("start_minute")
    mapping=[]
    for row in range(2,sheet.max_row+1):
        label=sheet.cell(row,1).value
        left,right=label.split("-")
        a,b=time_minutes(left),time_minutes(right)
        if b-a!=10: raise ValueError(f"Unexpected template interval: {label}")
        item=by_start.loc[a % 1440]
        sheet.cell(row,2,float(item.planned_grid_kwh))
        mapping.append({"excel_row":row,"template_interval":label,
                        "canonical_interval":item.interval,"source_excel_row":int(item.source_excel_row),
                        "cyclic_next_day":a>=1440,"grid_kwh":float(item.planned_grid_kwh)})
    _,blocks=reporting_tables(schedule)
    sheet=wb["充放电量"]
    for row,item in enumerate(blocks.itertuples(),2):
        sheet.cell(row,2,float(item.charge_kwh));sheet.cell(row,3,float(item.discharge_kwh))
    sheet.cell(2,5,float(schedule.soc_start_kwh.iloc[0]))
    sheet.cell(3,5,float(schedule.soc_end_kwh.iloc[-1]))
    target.parent.mkdir(parents=True,exist_ok=True)
    wb.save(target);wb.close()
    return pd.DataFrame(mapping)


def validate_problem1_workbook(target: Path, schedule, tolerance=2e-5):
    """Read-back numeric values and verify every unchanged template cell/style."""
    import numpy as np
    from openpyxl import load_workbook
    from .data import time_minutes
    from .models.evaluation import reporting_tables
    original=load_workbook(PATHS.result_template_dir / "result1.xlsx")
    actual=load_workbook(target)
    try:
        if original.sheetnames!=actual.sheetnames: raise ValueError("Sheet names changed")
        allowed={"计划购电量":{(r,2) for r in range(2,146)},
                 "充放电量":{(r,c) for r in range(2,8) for c in (2,3)}|{(2,5),(3,5)}}
        for so,sa in zip(original,actual):
            if (so.max_row,so.max_column)!=(sa.max_row,sa.max_column): raise ValueError("Template shape changed")
            if str(so.merged_cells)!=str(sa.merged_cells): raise ValueError("Merged cells changed")
            for row in so:
                for cell in row:
                    other=sa.cell(cell.row,cell.column)
                    # Writing a number materializes the default style of an empty cell.
                    if tuple(cell._style or [0]*9)!=tuple(other._style or [0]*9):
                        raise ValueError(f"Style changed: {cell.coordinate}")
                    if (cell.row,cell.column) not in allowed[so.title] and cell.value!=other.value:
                        raise ValueError(f"Template label changed: {cell.coordinate}")
        sheet=actual["计划购电量"]
        values=[]
        by_start=schedule.set_index("start_minute")
        for row in range(2,146):
            minute=time_minutes(sheet.cell(row,1).value.split("-")[0]) % 1440
            value=sheet.cell(row,2).value
            expected=float(by_start.loc[minute,"planned_grid_kwh"])
            if not isinstance(value,(float,int)) or abs(value-expected)>tolerance:
                raise ValueError(f"Wrong purchase cell: B{row}")
            values.append(value)
        _,blocks=reporting_tables(schedule)
        sheet=actual["充放电量"]
        for row,item in enumerate(blocks.itertuples(),2):
            if abs(sheet.cell(row,2).value-item.charge_kwh)>tolerance or abs(sheet.cell(row,3).value-item.discharge_kwh)>tolerance:
                raise ValueError("Storage aggregation mismatch")
        if abs(sheet.cell(2,5).value-schedule.soc_start_kwh.iloc[0])>tolerance or abs(sheet.cell(3,5).value-schedule.soc_end_kwh.iloc[-1])>tolerance:
            raise ValueError("SOC export mismatch")
        return {"passed":True,"template_labels_and_styles_preserved":True,
                "grid_sum_error_kwh":float(abs(np.sum(values)-schedule.planned_grid_kwh.sum())),
                "last_row_interpretation":"next-day 00:00--00:10 repeats cyclic first interval"}
    finally:
        original.close();actual.close()
