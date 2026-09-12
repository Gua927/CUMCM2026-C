"""Explicit, user-approved calendar interpretation of immutable template labels.

Official last-column text is preserved, but represents the current day's missing
00:00 slot, never the next day's contract. Header and per-date comments disclose
this exception. Internal observations use a separate chronological mapping.
"""
import numpy as np
import pandas as pd
from openpyxl.comments import Comment

from .data import time_minutes

NOTE = ("保留官方模板表头。经明确约定，本列存放本行日期当日00:00–00:10的计划量，"
        "不是次日购电量；当日汇总包含本列且不重复包含次日首段。实际交付日期见单元格批注及template_mapping.csv。")


def purchase_slot(label):
    if label == "0:00-0:10+1":
        return 0
    start, end = label.split("-")
    a, b = time_minutes(start), time_minutes(end)
    if a not in range(10, 1440, 10) or b != a+10:
        raise ValueError(f"Unexpected purchase template interval: {label}")
    return a//10


def annotate_calendar(wb, frame, titles):
    days = sorted(frame.date.unique())
    if frame.duplicated(["date", "slot"]).any():
        raise ValueError("Duplicate calendar delivery slots")
    convention = "daily_cyclic_start"
    if "time_convention" in frame:
        conventions = frame.time_convention.unique()
        if len(conventions) != 1:
            raise ValueError("Mixed time conventions in workbook")
        convention = conventions[0]
    mapping = []
    for title in titles:
        sheet = wb[title]
        sheet.cell(1, 145).comment = Comment(NOTE, "时间口径说明")
        for row in range(2, sheet.max_row+1):
            date = pd.Timestamp(sheet.cell(row, 1).value).normalize()
            text = str(date.date())
            if text not in days:
                raise ValueError(f"Missing calendar day {text}")
            sheet.cell(row, 145).comment = Comment(f"实际交付：{text} 00:00–00:10。"+NOTE, "时间口径说明")
            sheet.cell(row, 146).comment = Comment("本行日期00:00–24:00共144时段合计；末列按当日零点补位计入一次。", "时间口径说明")
            if title != titles[0]:
                continue
            for col in range(2, 146):
                label = sheet.cell(1, col).value
                slot = purchase_slot(label)
                start = date+pd.Timedelta(minutes=10*slot)
                source_date = date-pd.Timedelta(days=1) if slot == 0 and convention == "chronological_start" else date
                mapping.append({"row": row, "column": col, "issue_time": str(date),
                                "template_date": text, "template_label": label, "slot": slot,
                                "delivery_start": str(start), "delivery_end": str(start+pd.Timedelta(minutes=10)),
                                "canonical_interval": f"{slot//6:02d}:{slot%6*10:02d}-{(slot+1)//6:02d}:{(slot+1)%6*10:02d}",
                                "source_observation_date": str(source_date.date()),
                                "source_time": "0:00+1" if slot == 0 else f"{slot//6:02d}:{slot%6*10:02d}",
                                "time_convention": convention, "midnight_template_exception": slot == 0})
    return pd.DataFrame(mapping)


def audit_calendar_comments(wb, titles):
    """Independent audit of the unchanged template's exception and row sums."""
    for title in titles:
        sheet = wb[title]
        if sheet.cell(1, 145).value != "0:00-0:10+1":
            raise ValueError("Original midnight header changed")
        comment = sheet.cell(1, 145).comment
        if comment is None or comment.text != NOTE:
            raise ValueError("Missing explicit midnight interpretation")
        for row in range(2, sheet.max_row+1):
            date = str(pd.Timestamp(sheet.cell(row, 1).value).date())
            comment = sheet.cell(row, 145).comment
            if comment is None or f"实际交付：{date} 00:00–00:10。" not in comment.text:
                raise ValueError("Midnight delivery date not disclosed")
            values = np.asarray([sheet.cell(row, c).value for c in range(2, 146)], float)
            if not np.isfinite(values).all() or abs(values.sum()-sheet.cell(row, 146).value) > 2e-5:
                raise ValueError("Purchase day total does not equal all 144 displayed slots")
